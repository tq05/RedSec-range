#!/usr/bin/env python3
"""Build the defensive RedSec learning-guide PDF and PowerPoint deck."""
from __future__ import annotations

import hashlib
import os
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

# Homebrew installs Pango/GLib outside the default macOS dynamic-loader path.
# WeasyPrint loads those libraries during import, so set the fallback first.
if Path("/opt/homebrew/lib").is_dir():
    os.environ.setdefault("DYLD_FALLBACK_LIBRARY_PATH", "/opt/homebrew/lib")

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from weasyprint import HTML

ROOT = Path(__file__).resolve().parents[1]
TRAINING = ROOT / "docs" / "training"
GUIDE_MD = TRAINING / "LEARNING-GUIDE.md"
DECK_MD = TRAINING / "LEARNING-DECK.md"
GUIDE_PDF = TRAINING / "RedSec-Range-Learning-Guide.pdf"
DECK_PPTX = TRAINING / "RedSec-Range-Learning-Deck.pptx"
THEME = TRAINING / "learning-theme.css"
MANIFEST = TRAINING / "ARTIFACTS.sha256"

NAVY = RGBColor(7, 24, 39)
NAVY_2 = RGBColor(13, 39, 58)
WHITE = RGBColor(246, 249, 248)
MUTED = RGBColor(164, 183, 190)
CYAN = RGBColor(0, 196, 190)
RED = RGBColor(238, 83, 105)
ORANGE = RGBColor(241, 155, 78)
PURPLE = RGBColor(151, 121, 224)
FONT = "Avenir Next"


def run(command: list[str]) -> None:
    subprocess.run(command, cwd=ROOT, check=True)


def build_pdf() -> None:
    pandoc = shutil.which("pandoc")
    if not pandoc:
        raise SystemExit("pandoc is required to build the guide PDF")
    with tempfile.TemporaryDirectory(prefix="redsec-learning-") as temp:
        html = Path(temp) / "guide.html"
        css = Path(temp) / "learning-theme.css"
        css.write_text(THEME.read_text())
        run([
            pandoc,
            str(GUIDE_MD),
            "--from=gfm+attributes",
            "--to=html5",
            "--standalone",
            "--section-divs",
            "--toc",
            "--toc-depth=2",
            "--number-sections",
            f"--css={css.name}",
            "--metadata=lang:en-GB",
            f"--output={html}",
        ])
        fingerprint = f"source-sha256:{sha256(GUIDE_MD)}"
        html_text = html.read_text()
        html_text = html_text.replace(
            "</head>",
            f'<meta name="keywords" content="{fingerprint}">\n</head>',
        )
        html.write_text(html_text)
        HTML(filename=str(html), base_url=str(html.parent)).write_pdf(str(GUIDE_PDF))


def _comment(block: str, name: str) -> str:
    match = re.search(rf"^<!--\s*{re.escape(name)}:\s*(.+?)\s*-->$", block, re.M)
    return match.group(1).strip() if match else ""


def parse_deck() -> list[dict[str, object]]:
    blocks = re.split(r"\n\s*---\s*\n", DECK_MD.read_text().strip())
    slides: list[dict[str, object]] = []
    for block in blocks:
        title_match = re.search(r"^#\s+(.+)$", block, re.M)
        if not title_match:
            continue
        takeaway_match = re.search(r"^>\s*(.+)$", block, re.M)
        rows = []
        for bullet in re.findall(r"^-\s+(.+)$", block, re.M):
            labelled = re.match(r"\*\*([^*]+?):\*\*\s*(.+)", bullet.strip())
            if labelled:
                rows.append({"label": labelled.group(1), "text": labelled.group(2)})
            else:
                rows.append({"label": "NOTE", "text": bullet.strip()})
        slides.append({
            "title": title_match.group(1).strip(),
            "takeaway": takeaway_match.group(1).strip() if takeaway_match else "",
            "kind": _comment(block, "kind"),
            "key": _comment(block, "key"),
            "range": _comment(block, "range"),
            "pair": _comment(block, "pair"),
            "rows": rows,
        })
    return slides


def set_fill(shape, color: RGBColor, transparency: int = 0) -> None:
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if transparency:
        shape.fill.transparency = transparency
    shape.line.fill.background()


def add_text(slide, text: str, x: float, y: float, w: float, h: float,
             size: float, color: RGBColor, *, bold: bool = False,
             font: str = FONT, align=PP_ALIGN.LEFT,
             valign=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run_ = paragraph.add_run()
    run_.text = text
    run_.font.name = font
    run_.font.size = Pt(size)
    run_.font.bold = bold
    run_.font.color.rgb = color
    return box


def add_background(slide, accent: RGBColor, index: int, total: int,
                   slide_data: dict[str, object]) -> None:
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = NAVY

    band = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(.14), Inches(7.5))
    set_fill(band, accent)

    # restrained technical grid
    for x in (9.25, 10.25, 11.25, 12.25):
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(x), Inches(.35), Inches(.012), Inches(6.55)
        )
        set_fill(line, RGBColor(27, 55, 72))
    for y in (1.35, 2.35, 3.35, 4.35, 5.35, 6.35):
        line = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE, Inches(8.8), Inches(y), Inches(4.1), Inches(.012)
        )
        set_fill(line, RGBColor(27, 55, 72))

    kind = str(slide_data["kind"]).upper()
    add_text(slide, "REDSEC RANGE  //  DEFENSIVE LEARNING EDITION", .55, .25, 6.5, .3,
             8.5, MUTED, bold=True)
    add_text(slide, f"{index:03d} / {total}", 11.55, .22, 1.05, .35, 8.5, accent,
             bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, f"{slide_data['range']}  •  {kind}", .55, 6.76, 4.2, .28,
             7.5, accent, bold=True)
    add_text(slide, f"PAIR {slide_data['pair']}", 10.85, 6.76, 1.8, .28,
             7.5, MUTED, bold=True, align=PP_ALIGN.RIGHT)

    progress_bg = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(.55), Inches(7.17), Inches(12.1), Inches(.035)
    )
    set_fill(progress_bg, RGBColor(34, 60, 76))
    progress = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(.55), Inches(7.17),
        Inches(12.1 * index / total), Inches(.035)
    )
    set_fill(progress, accent)


def accent_for(range_id: str) -> RGBColor:
    return {
        "builtin": CYAN,
        "dvaa": PURPLE,
        "mcpgoat": ORANGE,
        "dvmcp": RGBColor(55, 211, 176),
        "appsecco-mcp": RED,
    }.get(range_id, CYAN)


def add_content_slide(
    prs: Presentation,
    slide_data: dict[str, object],
    index: int,
    total: int,
) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    accent = accent_for(str(slide_data["range"]))
    add_background(slide, accent, index, total, slide_data)

    title = str(slide_data["title"])
    title_size = 27 if len(title) < 62 else 23
    add_text(slide, title, .72, .76, 11.35, .92, title_size, WHITE, bold=True)

    takeaway = str(slide_data["takeaway"])
    box = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.72), Inches(1.78), Inches(11.9), Inches(.82)
    )
    set_fill(box, NAVY_2)
    box.line.color.rgb = accent
    add_text(slide, takeaway, 1.0, 1.94, 11.35, .48, 15.5, accent, bold=True,
             valign=MSO_ANCHOR.MIDDLE)

    rows = list(slide_data["rows"])
    for position, row in enumerate(rows[:4]):
        y = 2.84 + position * .91
        panel = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE, Inches(.72), Inches(y),
            Inches(11.9), Inches(.76)
        )
        set_fill(panel, RGBColor(10, 31, 47))
        panel.line.color.rgb = RGBColor(31, 64, 80)
        label = str(row["label"])
        body = str(row["text"])
        add_text(slide, label, .93, y + .20, 1.34, .30, 8.5, accent, bold=True,
                 valign=MSO_ANCHOR.MIDDLE)
        size = 14.0 if len(body) < 150 else 12.4 if len(body) < 230 else 10.8
        add_text(slide, body, 2.28, y + .10, 10.02, .56, size, WHITE,
                 valign=MSO_ANCHOR.MIDDLE)


def build_deck() -> None:
    slides = parse_deck()
    if len(slides) != 220:
        raise SystemExit(f"expected 220 paired challenge slides, found {len(slides)}")
    for index in range(0, len(slides), 2):
        challenge, resolution = slides[index:index + 2]
        if challenge["kind"] != "challenge" or resolution["kind"] != "resolution":
            raise SystemExit(f"invalid pair order at slides {index + 1}-{index + 2}")
        if challenge["key"] != resolution["key"]:
            raise SystemExit(f"pair key mismatch at slides {index + 1}-{index + 2}")
    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "RedSec Range — Defensive Challenge Learning Deck"
    prs.core_properties.subject = "Defensive review and remediation for AI, A2A, and MCP challenge concepts"
    prs.core_properties.author = "RedSec Range"
    prs.core_properties.keywords = "AI security, MCP, A2A, defensive review, training"
    prs.core_properties.comments = f"source-sha256:{sha256(DECK_MD)}"

    for index, slide_data in enumerate(slides, start=1):
        add_content_slide(prs, slide_data, index, len(slides))
    prs.save(DECK_PPTX)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_manifest() -> None:
    paths = (GUIDE_MD, DECK_MD, THEME, GUIDE_PDF, DECK_PPTX)
    lines = ["# sha256  path", *[f"{sha256(path)}  {path.relative_to(ROOT)}" for path in paths]]
    MANIFEST.write_text("\n".join(lines) + "\n")


def main() -> int:
    TRAINING.mkdir(parents=True, exist_ok=True)
    build_pdf()
    build_deck()
    write_manifest()
    print(f"wrote {GUIDE_PDF}")
    print(f"wrote {DECK_PPTX}")
    print(f"wrote {MANIFEST}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
